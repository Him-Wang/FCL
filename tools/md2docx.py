# -*- coding: utf-8 -*-
"""Markdown -> DOCX 高保真转换（保留标题层级 / 表格 / 代码块 / 列表 / 引用 / 行内格式）"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BRAND = '4D6BFE'        # DeepSeek 品牌蓝，用作标题与表头
GREY = '5F5E5A'
CODE_BG = 'F4F5F8'
CN = '微软雅黑'
CN_MONO = '等线'

INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def _rfonts(run, ascii_name, eastasia):
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.append(rf)
    rf.set(qn('w:ascii'), ascii_name)
    rf.set(qn('w:hAnsi'), ascii_name)
    rf.set(qn('w:eastAsia'), eastasia)


def style_run(run, size=10.5, bold=False, color=None, mono=False, italic=False):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if mono:
        run.font.name = 'Consolas'
        _rfonts(run, 'Consolas', CN_MONO)
    else:
        run.font.name = CN
        _rfonts(run, CN, CN)
    return run


def shade(el, fill):
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), fill)
    el.append(sh)


def add_inline(p, text, size=10.5, color=None, base_bold=False):
    """解析 **bold** / `code` / [link](url) 并写入段落"""
    for seg in INLINE.split(text):
        if not seg:
            continue
        if seg.startswith('**') and seg.endswith('**') and len(seg) > 4:
            style_run(p.add_run(seg[2:-2]), size, True, color)
        elif seg.startswith('`') and seg.endswith('`') and len(seg) > 2:
            style_run(p.add_run(seg[1:-1]), size - 0.5, base_bold, 'A32D2D', mono=True)
        elif seg.startswith('[') and '](' in seg and seg.endswith(')'):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', seg)
            if m:
                r = style_run(p.add_run(m.group(1)), size, base_bold, '185FA5')
                r.font.underline = True
            else:
                style_run(p.add_run(seg), size, base_bold, color)
        else:
            style_run(p.add_run(seg), size, base_bold, color)


def add_heading(doc, text, level):
    sizes = {1: 20, 2: 15, 3: 12.5, 4: 11.5}
    colors = {1: BRAND, 2: '2C2C2A', 3: '3C3489', 4: '5F5E5A'}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level <= 2 else 10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    style_run(p.add_run(text), sizes.get(level, 11), True, colors.get(level, '000000'))
    if level == 1:
        pr = p._p.get_or_add_pPr()
        pbdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '4')
        bottom.set(qn('w:color'), BRAND)
        pbdr.append(bottom)
        pr.append(pbdr)
    return p


def add_code_block(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        shade(p._p.get_or_add_pPr(), CODE_BG)
        style_run(p.add_run(line if line else ' '), 9, mono=True, color='1F2328')
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_table(doc, rows):
    if not rows:
        return
    ncol = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci in range(ncol):
            cell = t.cell(ri, ci)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, str(row[ci] if ci < len(row) else ''),
                       size=9.5, base_bold=(ri == 0))
            if ri == 0:
                shade(cell._tc.get_or_add_tcPr(), 'EEEDFE')
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def parse_row(line):
    s = line.strip()
    if s.startswith('|'):
        s = s[1:]
    if s.endswith('|'):
        s = s[:-1]
    return [c.strip() for c in s.split('|')]


def is_sep(line):
    return bool(re.match(r'^\|?[\s:\-\|]+\|?$', line.strip())) and '-' in line


def convert(md_path, docx_path, title=None, subtitle=None):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)

    normal = doc.styles['Normal']
    normal.font.name = CN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), CN)

    # 封面标题
    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(90)
        style_run(p.add_run(title), 26, True, BRAND)
        if subtitle:
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_before = Pt(10)
            style_run(p2.add_run(subtitle), 12, False, GREY)
        doc.add_page_break()

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()

        # 代码块
        if s.startswith('```'):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, buf)
            continue

        if not s:
            i += 1
            continue

        # 分隔线
        if re.match(r'^-{3,}$', s) or re.match(r'^\*{3,}$', s):
            i += 1
            continue

        # 标题
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            add_heading(doc, m.group(2).strip(), min(len(m.group(1)), 4))
            i += 1
            continue

        # 表格
        if s.startswith('|') and i + 1 < n and is_sep(lines[i + 1]):
            rows = [parse_row(s)]
            i += 2
            while i < n and lines[i].strip().startswith('|'):
                rows.append(parse_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # 引用
        if s.startswith('>'):
            txt = s.lstrip('>').strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            pr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            left = OxmlElement('w:left')
            left.set(qn('w:val'), 'single')
            left.set(qn('w:sz'), '12')
            left.set(qn('w:space'), '8')
            left.set(qn('w:color'), BRAND)
            pbdr.append(left)
            pr.append(pbdr)
            shade(pr, 'F7F8FC')
            add_inline(p, txt, size=10, color=GREY)
            i += 1
            continue

        # 列表
        m = re.match(r'^([-*])\s+(.*)$', s)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(2), size=10.5)
            i += 1
            continue
        m = re.match(r'^(\d+)\.\s+(.*)$', s)
        if m:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(2), size=10.5)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.4
        add_inline(p, s, size=10.5)
        i += 1

    doc.save(docx_path)
    print('saved: ' + docx_path)


def merge(md_paths, drop_first_h1=True):
    """拼接多个 md；可选丢弃首个 h1（封面已有标题）"""
    parts = []
    for k, p in enumerate(md_paths):
        with open(p, 'r', encoding='utf-8') as f:
            txt = f.read()
        if k == 0 and drop_first_h1:
            txt = re.sub(r'^#\s+.*\n', '', txt, count=1)
        parts.append(txt.strip())
    return '\n\n'.join(parts) + '\n'


if __name__ == '__main__':
    # 用法: md2docx.py <out.docx> <title> <subtitle> <in1.md> [in2.md ...]
    out = sys.argv[1]
    title = sys.argv[2] if len(sys.argv) > 2 else None
    subtitle = sys.argv[3] if len(sys.argv) > 3 else None
    inputs = sys.argv[4:]
    if not inputs:
        print('usage: md2docx.py out.docx "title" "subtitle" in1.md [in2.md ...]')
        sys.exit(1)
    tmp = out + '.merged.md'
    with open(tmp, 'w', encoding='utf-8') as f:
        f.write(merge(inputs))
    convert(tmp, out, title=title, subtitle=subtitle)
    import os
    os.remove(tmp)
